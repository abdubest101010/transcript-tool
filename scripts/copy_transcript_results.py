import sys
import os
import re
from docx import Document

def normalize(text):
    return re.sub(r'[^a-zA-Z0-9]', '', (text or '').lower())

def extract_results(doc_path):
    """
    Extracts subject marks, overall marks (Total, Average), and Rank from document tables.
    """
    doc = Document(doc_path)
    extracted_data = {}
    
    # Locate the table with Subject results
    for table in doc.tables:
        has_subjects = False
        for row in table.rows:
            row_text = " ".join(cell.text for cell in row.cells)
            if "subject" in row_text.lower() or "grade" in row_text.lower():
                has_subjects = True
                break
        
        if has_subjects:
            for row in table.rows:
                if len(row.cells) < 2:
                    continue
                first_col = row.cells[0].text.strip()
                norm_name = normalize(first_col)
                if not norm_name:
                    continue
                
                # Extract score cell values (cells 1 to end)
                scores = [c.text.strip() for c in row.cells[1:]]
                extracted_data[norm_name] = {
                    'original_name': first_col,
                    'scores': scores
                }
    return extracted_data

def copy_results(source_doc_path, target_doc_path, output_doc_path):
    """
    Copies Subject marks, Total, Average, and Rank from source to target document.
    Leaves 100% of target document formatting, styles, images, and other text intact.
    """
    source_results = extract_results(source_doc_path)
    target_doc = Document(target_doc_path)
    
    for table in target_doc.tables:
        has_subjects = False
        for row in table.rows:
            row_text = " ".join(cell.text for cell in row.cells)
            if "subject" in row_text.lower() or "grade" in row_text.lower():
                has_subjects = True
                break
        
        if has_subjects:
            for row in table.rows:
                if len(row.cells) < 2:
                    continue
                
                first_col = row.cells[0].text.strip()
                norm_name = normalize(first_col)
                
                # Check if this row is a subject, total, average, or rank
                if norm_name in source_results:
                    src_scores = source_results[norm_name]['scores']
                    
                    # Update each score cell (index 1 to end)
                    for col_idx, cell in enumerate(row.cells[1:]):
                        if col_idx < len(src_scores):
                            new_val = src_scores[col_idx]
                            
                            # Preserve formatting by replacing only text inside existing paragraphs/runs
                            if cell.paragraphs:
                                p = cell.paragraphs[0]
                                if p.runs:
                                    p.runs[0].text = new_val
                                    for r in p.runs[1:]:
                                        r.text = ""
                                else:
                                    p.text = new_val
                            else:
                                cell.text = new_val
                                
    target_doc.save(output_doc_path)
    print(f"Successfully saved updated document to: {output_doc_path}")

if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "sample/Frtuna Transcript-1.docx"
    tgt = sys.argv[2] if len(sys.argv) > 2 else "sample/Yonatan Transcript.docx"
    out = sys.argv[3] if len(sys.argv) > 3 else "public/Updated_Student_Transcript.docx"
    copy_results(src, tgt, out)
